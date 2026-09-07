"""
Fresher.AI — Master Knowledge Base Excel Generator
Builds a production-ready, 17-sheet Excel workbook with openpyxl.
Adheres strictly to the JSON-first data model:
- No merged cells in data tables
- Clean header row with styling
- Auto-adjusted column widths
- Stable snake_case IDs and verified references
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from data import SHEETS_REGISTRY

# Styling configuration
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")  # Deep Navy Blue
DATA_FONT = Font(name="Calibri", size=10, color="000000")
BORDER_THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB")
)

def build_workbook(output_path: str):
    print("=" * 60)
    print("Fresher.AI — Generating Career RAG Knowledge Base Spreadsheet")
    print("=" * 60)
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    total_rows = 0
    
    for idx, sheet_meta in enumerate(SHEETS_REGISTRY, 1):
        sheet_name = sheet_meta["name"]
        data_fn = sheet_meta["data_fn"]
        headers_fn = sheet_meta["headers_fn"]
        
        headers = headers_fn()
        rows = data_fn()
        
        ws = wb.create_sheet(title=sheet_name)
        
        # 1. Write Header Row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER_THIN
        
        ws.row_dimensions[1].height = 28
        
        # 2. Write Data Rows
        for row_num, item in enumerate(rows, 2):
            for col_num, header in enumerate(headers, 1):
                val = item.get(header, "")
                # Convert bools to explicit True/False strings or booleans
                if isinstance(val, bool):
                    val = "TRUE" if val else "FALSE"
                elif isinstance(val, (list, tuple)):
                    val = ";".join(str(x) for x in val)
                
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER_THIN
                
                # Align numbers right, text left
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif header == "embedding_text" or "description" in header or "breakdown" in header or "outcomes" in header or "tasks" in header or "topics" in header or "evidence" in header or "deliverable" in header or "goal" in header:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top")
            
            ws.row_dimensions[row_num].height = 20
        
        # 3. Freeze Header Row
        ws.freeze_panes = "A2"
        
        # 4. Auto-calculate column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            header_name = col[0].value or ""
            
            if header_name == "embedding_text":
                ws.column_dimensions[col_letter].width = 45
            elif any(k in header_name for k in ["description", "what_to_learn", "deliverable", "outcome", "evidence", "task", "goal"]):
                ws.column_dimensions[col_letter].width = 38
            elif "answer" in header_name or "breakdown" in header_name or "topic" in header_name:
                ws.column_dimensions[col_letter].width = 35
            else:
                max_len = max(len(str(cell.value or '')) for cell in col[:15])  # Check top 15 rows for speed
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 32)
        
        sheet_row_count = len(rows)
        total_rows += sheet_row_count
        print(f"[{idx:02d}/17] Created Sheet '{sheet_name:<20}' -> {sheet_row_count:>3} records, {len(headers):>2} columns")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    try:
        wb.save(output_path)
        print("=" * 60)
        print(f"SUCCESS: Saved Knowledge Base Workbook to:\n{output_path}")
    except PermissionError:
        fallback_path = output_path.replace(".xlsx", "_Master.xlsx")
        wb.save(fallback_path)
        print("=" * 60)
        print(f"NOTE: Target Excel file was open in Excel. Saved updated Knowledge Base to:\n{fallback_path}")

    print(f"Total Sheets: {len(SHEETS_REGISTRY)} | Total Entities / Rows: {total_rows}")
    print("=" * 60)

if __name__ == "__main__":
    output_excel = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Fresher_AI_Career_Knowledge_Base.xlsx")
    build_workbook(output_excel)

